"""テスト定義 → Excel テスト仕様書 生成ツール.

YAML または プレーンテキスト(.txt) でテストケースを定義し、Excel仕様書に変換する。
Excelを手書きする代わりに定義ファイルで記述でき、Git差分やIDEの補完が使える。

プレーンテキスト形式:
    1行1ステップの簡易記法。YAML 構文を一切使わずにテストを定義できる。
    python gen_spec.py test.txt -o spec.xlsx

データ駆動テスト:
    data_source でCSV/JSONファイルを指定すると、データ行ごとにステップを展開する。
    ステップ内で ${data:列名} を使ってデータ値を参照できる。

使い方:
    python gen_spec.py test_spec.yaml -o spec.xlsx
    python gen_spec.py test_spec.yaml              # → test_spec.xlsx に出力
    python gen_spec.py test.txt -o spec.xlsx       # プレーンテキスト入力
"""

import argparse
import csv as csv_mod
import json
import re
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# スタイル定数
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(name="Yu Gothic", size=10, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Yu Gothic", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

# --- 自然言語パターンマッチ ---

# 操作パターン（優先度順）
NL_ACTION_PATTERNS = [
    # 「セレクタ」に「値」を入力
    (r"「(.+?)」に「(.+?)」を入力", {"action": "input", "selector": 1, "input": 2}),
    # 〇〇欄に「値」を入力
    (r"(.+?)欄に「(.+?)」を入力", {"action": "input", "selector": 1, "input": 2}),
    # 「セレクタ」で「値」を選択
    (r"「(.+?)」で「(.+?)」を選択", {"action": "select", "selector": 1, "input": 2}),
    # 「セレクタ」をクリック
    (r"「(.+?)」をクリック", {"action": "click", "selector": 1}),
    # 〇〇ボタンをクリック
    (r"(.+?)ボタンをクリック", {"action": "click", "selector": 1}),
    # 〇〇リンクをクリック
    (r"(.+?)リンクをクリック", {"action": "click", "selector": 1}),
    # 〇〇に遷移 / 〇〇を開く / 〇〇へ遷移 / 〇〇にアクセス
    (r"「?(.+?)」?(?:に遷移|を開く|へ遷移|にアクセス)", {"action": "navigate", "input": 1}),
    # N秒待機
    (r"(\d+)秒待機", {"action": "wait", "input_sec": 1}),
    # 「セレクタ」が表示されるまで待機
    (r"「(.+?)」が表示されるまで待機", {"action": "wait_for", "selector": 1}),
    # 「セレクタ」にマウスオーバー / ホバー
    (r"「(.+?)」に(?:マウスオーバー|ホバー)", {"action": "hover", "selector": 1}),
    # Enterキーを押す / Tabキーを押す / Escapeキーを押す
    (r"(Enter|Tab|Escape)キーを押す", {"action": "keyboard", "input": 1}),
    # 「セレクタ」の値を保存 / キャプチャ / 取得
    (r"「(.+?)」の値を(?:保存|キャプチャ|取得)", {"action": "capture", "selector": 1}),
    # 「ファイル」をアップロード
    (r"「(.+?)」をアップロード", {"action": "upload", "input": 1}),
    # 「セレクタ」までスクロール
    (r"「(.+?)」までスクロール", {"action": "scroll", "selector": 1}),
]

# 検証パターン
NL_VERIFY_PATTERNS = [
    # 「セレクタ」に「値」と表示
    (r"「(.+?)」に「(.+?)」と表示", {"verify": "text", "target": 1, "expected": 2}),
    # 「セレクタ」に「値」を含む
    (r"「(.+?)」に「(.+?)」を含む", {"verify": "text", "target": 1, "expected_prefix": "contains:", "expected": 2}),
    # 「セレクタ」が表示されること / 「セレクタ」が表示される
    (r"「(.+?)」が表示される(?:こと)?", {"verify": "visible", "target": 1}),
    # 「セレクタ」が非表示
    (r"「(.+?)」が非表示", {"verify": "hidden", "target": 1}),
    # URLが「値」であること
    (r"URLが「(.+?)」", {"verify": "url", "expected": 1}),
    # スクリーンショットを取得 / スクリーンショット
    (r"スクリーンショット(?:を取得)?", {"verify": "screenshot"}),
]


def _interpret_natural_language(step_def: dict, selector_aliases: dict | None = None) -> dict:
    """テスト項目(item)の自然言語記述を構造化フィールドに変換する.

    action / verify_type が既に指定されている場合はスキップ（明示指定を優先）。

    Args:
        step_def: ステップ定義の辞書
        selector_aliases: 日本語名→CSSセレクタの辞書（任意）

    Returns:
        構造化フィールドが補完されたステップ定義
    """
    if step_def.get("action") or step_def.get("verify") or step_def.get("verify_type"):
        return step_def

    text = step_def.get("item", "").strip()
    if not text:
        return step_def

    aliases = selector_aliases or {}
    result = dict(step_def)

    def _resolve_selector(raw: str) -> str:
        """セレクタ辞書で日本語名をCSSセレクタに解決する."""
        return aliases.get(raw, raw)

    # 操作パターンを試行
    for pattern, mapping in NL_ACTION_PATTERNS:
        m = re.search(pattern, text)
        if m:
            for key, val in mapping.items():
                if key == "input_sec":
                    result["input"] = str(int(m.group(val)) * 1000)
                elif isinstance(val, int):
                    captured = m.group(val)
                    if key == "selector":
                        captured = _resolve_selector(captured)
                    result[key] = captured
                else:
                    result[key] = val
            return result

    # 検証パターンを試行
    for pattern, mapping in NL_VERIFY_PATTERNS:
        m = re.search(pattern, text)
        if m:
            prefix = mapping.get("expected_prefix", "")
            for key, val in mapping.items():
                if key == "expected_prefix":
                    continue
                if isinstance(val, int):
                    captured = m.group(val)
                    if key == "target":
                        captured = _resolve_selector(captured)
                    if key == "expected" and prefix:
                        result[key] = prefix + captured
                    else:
                        result[key] = captured
                else:
                    result[key] = val
            return result

    # パターン不一致 → 警告チェック
    warning = _check_nl_warning(text)
    if warning:
        result["_nl_warning"] = warning
    return result


# --- NL 警告検出 ---

NL_HINT_KEYWORDS_ACTION = (
    "クリック", "入力", "選択", "遷移", "開く", "アクセス",
    "待機", "ホバー", "マウスオーバー", "アップロード",
    "スクロール", "キャプチャ", "保存", "取得", "押す",
)
NL_HINT_KEYWORDS_VERIFY = ("表示", "非表示", "含む", "スクリーンショット")


def _check_nl_warning(text: str) -> str | None:
    """NLパターン不一致時に、書き間違いの可能性を警告する."""
    # 括弧の不整合
    open_count = text.count("「")
    close_count = text.count("」")
    if open_count != close_count:
        return (
            f"括弧「」の数が不一致です（開: {open_count}, 閉: {close_count}）。"
            " 正しい記法例: 「#selector」をクリック"
        )

    # 操作キーワード
    for kw in NL_HINT_KEYWORDS_ACTION:
        if kw in text:
            return (
                f"操作キーワード「{kw}」が見つかりましたが、"
                "パターンに一致しません。記法を確認してください。"
                " 例: 「#selector」をクリック、「#input」に「値」を入力"
            )

    # 検証キーワード
    for kw in NL_HINT_KEYWORDS_VERIFY:
        if kw in text:
            return (
                f"検証キーワード「{kw}」が見つかりましたが、"
                "パターンに一致しません。記法を確認してください。"
                " 例: 「#selector」が表示されること、「#msg」に「値」と表示"
            )

    return None


VALID_ACTIONS = {
    "navigate", "click", "input", "select", "wait", "wait_for",
    "upload", "hover", "scroll", "keyboard",
    "alert_accept", "alert_dismiss", "iframe", "include", "skip",
    "capture", "new_tab", "switch_tab", "close_tab", "download",
    "if_ok", "if_ng", "",
}
VALID_VERIFY_TYPES = {
    "text", "value", "visible", "hidden", "url", "screenshot", "db",
    "download", "",
}


def create_sheet(wb, sheet_name, test_cases, project_name="", is_first=False):
    """Excelシートにテストケースを書き込む."""
    if is_first:
        ws = wb.active
        ws.title = sheet_name
    else:
        ws = wb.create_sheet(sheet_name)

    ws["A1"] = "テスト仕様書"
    ws["A1"].font = Font(name="Yu Gothic", size=14, bold=True)
    ws["A2"] = "プロジェクト名:"
    ws["B2"] = project_name
    ws["A3"] = "テスト日:"
    ws["A4"] = "テスト担当者:"
    ws["B4"] = "自動テスト"

    headers = [
        ("A", "No.", 6), ("B", "テスト項目", 30), ("C", "操作種別", 12),
        ("D", "対象セレクタ", 30), ("E", "入力値", 25),
        ("F", "検証種別", 12), ("G", "検証対象", 35),
        ("H", "期待値", 30), ("I", "結果", 8),
        ("J", "エビデンス", 30), ("K", "備考", 25),
    ]
    for col, title, width in headers:
        cell = ws[f"{col}6"]
        cell.value = title
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        ws.column_dimensions[col].width = width

    for i, tc in enumerate(test_cases):
        row = 7 + i
        values = [
            tc.get("no", i + 1), tc.get("item", ""),
            tc.get("action", ""), tc.get("selector", ""),
            tc.get("input", ""), tc.get("verify_type", ""),
            tc.get("verify_target", ""), tc.get("expected", ""),
            "", "", tc.get("note", ""),
        ]
        for j, val in enumerate(values):
            col = chr(65 + j)
            cell = ws[f"{col}{row}"]
            cell.value = val
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            cell.alignment = LEFT if j > 0 else CENTER
        ws.row_dimensions[row].height = 80

    return ws


def normalize_step(step_def: dict, step_no: int,
                    selector_aliases: dict | None = None) -> dict:
    """YAML上の短縮記法を正規化する.

    短縮記法:
        - { action: click, selector: "#btn" }
        - { verify: text, target: "#msg", expected: "OK" }
        - { verify: db, sql: "SELECT ...", expected: "rowcount:3" }
        - { action: include, sheet: "共通ログイン" }

    自然言語:
        - { item: "「#btn」をクリック" }
        - { item: "「#msg」に「OK」と表示" }
    """
    # 自然言語解釈（action/verify が未指定の場合のみ）
    step_def = _interpret_natural_language(step_def, selector_aliases)

    tc = {"no": step_no}

    # item（テスト項目名）
    tc["item"] = step_def.get("item", "")

    # action 系
    action = step_def.get("action", "")
    tc["action"] = action
    tc["selector"] = step_def.get("selector", "")
    tc["input"] = step_def.get("input", "")

    # include の短縮記法: { action: include, sheet: "xxx" }
    if action == "include" and "sheet" in step_def:
        tc["input"] = step_def["sheet"]

    # verify 系: { verify: text, target: "#msg" } の短縮記法対応
    if "verify" in step_def:
        tc["verify_type"] = step_def["verify"]
        tc["verify_target"] = step_def.get(
            "target", step_def.get("sql",
                                   step_def.get("verify_target", "")))
    else:
        tc["verify_type"] = step_def.get("verify_type", "")
        tc["verify_target"] = step_def.get("verify_target", "")

    # db検証の短縮: { verify: db, sql: "SELECT ..." }
    if "sql" in step_def and not tc["verify_target"]:
        tc["verify_type"] = "db"
        tc["verify_target"] = step_def["sql"]

    tc["expected"] = step_def.get("expected", "")
    note = step_def.get("note", "")
    # retry / retry_wait を備考欄に埋め込む
    retry = step_def.get("retry")
    if retry:
        note = f"retry:{retry} {note}".strip()
    retry_wait = step_def.get("retry_wait")
    if retry_wait:
        note = f"retry_wait:{retry_wait} {note}".strip()
    tc["note"] = note

    # item が空の場合、action/verify から自動生成
    if not tc["item"]:
        if action == "include":
            tc["item"] = f"共通手順: {tc['input']}"
        elif action == "navigate":
            tc["item"] = f"ページを開く"
        elif action:
            tc["item"] = f"{action}: {tc['selector'] or tc['input']}"
        elif tc["verify_type"]:
            tc["item"] = f"検証: {tc['verify_type']}"

    return tc


def validate_steps(sheet_name: str, steps: list[dict]) -> list[str]:
    """テストステップのバリデーション."""
    errors = []
    for i, step in enumerate(steps):
        prefix = f"[{sheet_name}] Step {i + 1}"
        action = step.get("action", "")
        if action and action not in VALID_ACTIONS:
            errors.append(f"{prefix}: 無効な操作種別 '{action}'")

        verify = step.get("verify", step.get("verify_type", ""))
        if verify and verify not in VALID_VERIFY_TYPES:
            errors.append(f"{prefix}: 無効な検証種別 '{verify}'")

        if action in ("click", "input", "select", "hover", "wait_for",
                       "upload", "iframe") and not step.get("selector"):
            errors.append(f"{prefix}: action='{action}' にセレクタが未指定")

        if action == "navigate" and not step.get("input"):
            errors.append(f"{prefix}: action='navigate' にURLが未指定")

        if action == "include" and not step.get("input") and not step.get("sheet"):
            errors.append(f"{prefix}: action='include' にシート名が未指定")

        if action == "capture" and not step.get("selector"):
            errors.append(f"{prefix}: action='capture' にセレクタが未指定")
        if action == "capture" and not step.get("input"):
            errors.append(f"{prefix}: action='capture' に変数名(input)が未指定")

        if action == "switch_tab" and not step.get("input"):
            errors.append(f"{prefix}: action='switch_tab' にタブ名(input)が未指定")

        if action == "download" and not step.get("selector"):
            errors.append(f"{prefix}: action='download' にセレクタが未指定")

    return errors


def _expand_data_source(steps: list[dict], data_path: str,
                        yaml_dir: Path) -> list[dict]:
    """data_source で指定されたCSV/JSONを読み込み、ステップをデータ行分展開する.

    ステップ内の ${data:列名} をデータ値で置換する。

    Args:
        steps: 元のステップ定義リスト
        data_path: CSV/JSONファイルパス（YAML相対 or 絶対）
        yaml_dir: YAMLファイルのディレクトリ（相対パス解決用）

    Returns:
        展開されたステップリスト
    """
    path = Path(data_path)
    if not path.is_absolute():
        path = yaml_dir / path
    if not path.exists():
        print(f"エラー: data_source ファイルが見つかりません: {path}", file=sys.stderr)
        sys.exit(1)

    # データ読み込み
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as f:
            reader = csv_mod.DictReader(f)
            data_rows = list(reader)
    elif suffix == ".json":
        with open(path, encoding="utf-8") as f:
            data_rows = json.load(f)
            if not isinstance(data_rows, list):
                print(f"エラー: data_source JSON はリストである必要があります: {path}",
                      file=sys.stderr)
                sys.exit(1)
    else:
        print(f"エラー: data_source は .csv または .json のみ対応: {path}",
              file=sys.stderr)
        sys.exit(1)

    if not data_rows:
        return steps  # データなし → 元のステップをそのまま返す

    # ステップ展開: データ行ごとにステップを複製して ${data:col} を置換
    expanded = []
    for row_idx, row_data in enumerate(data_rows, 1):
        row_label = row_data.get("name", row_data.get("id", str(row_idx)))
        for step in steps:
            new_step = dict(step)
            # 各フィールドの ${data:col} を置換
            for key in ("selector", "input", "expected", "item",
                        "verify_target", "target", "sql", "note", "sheet"):
                if key in new_step and isinstance(new_step[key], str):
                    new_step[key] = re.sub(
                        r"\$\{data:(\w+)\}",
                        lambda m: str(row_data.get(m.group(1), m.group(0))),
                        new_step[key],
                    )
            # テスト項目名にデータ識別子を付与
            if "item" in new_step and new_step["item"]:
                new_step["item"] = f"[row{row_idx}:{row_label}] {new_step['item']}"
            expanded.append(new_step)

    return expanded


def gen_spec(input_path: str, output_path: str | None = None):
    """テスト定義（YAML / テキスト）からExcelテスト仕様書を生成する."""
    input_path = Path(input_path)
    if not input_path.exists():
        print(f"エラー: ファイルが見つかりません: {input_path}", file=sys.stderr)
        sys.exit(1)

    # プレーンテキスト形式の自動判定
    if input_path.suffix.lower() == ".txt":
        from text2spec import parse_text_spec
        spec = parse_text_spec(input_path)
    else:
        with open(input_path, encoding="utf-8") as f:
            spec = yaml.safe_load(f) or {}

    sheets = spec.get("sheets", [])
    if not sheets:
        print("エラー: 'sheets' が定義されていません", file=sys.stderr)
        sys.exit(1)

    project_name = spec.get("project", "")
    selector_aliases = spec.get("selectors", {})

    # NL解釈後にバリデーション（NLで補完されたフィールドも検証対象にする）
    all_errors = []
    all_warnings = []
    for sheet_def in sheets:
        name = sheet_def.get("name", "Sheet")
        steps = sheet_def.get("steps", [])
        interpreted_steps = [_interpret_natural_language(s, selector_aliases) for s in steps]
        all_errors.extend(validate_steps(name, interpreted_steps))
        # NL警告を収集
        for j, interp in enumerate(interpreted_steps):
            nl_warning = interp.get("_nl_warning")
            if nl_warning:
                all_warnings.append(f"[{name}] Step {j + 1}: {nl_warning}")

    if all_warnings:
        print("自然言語の警告:", file=sys.stderr)
        for warn in all_warnings:
            print(f"  警告 {warn}", file=sys.stderr)

    if all_errors:
        print("バリデーションエラー:", file=sys.stderr)
        for err in all_errors:
            print(f"  {err}", file=sys.stderr)
        sys.exit(1)

    # Excel生成
    input_dir = input_path.parent
    wb = Workbook()
    for i, sheet_def in enumerate(sheets):
        name = sheet_def.get("name", f"Sheet{i + 1}")
        steps = sheet_def.get("steps", [])
        # データ駆動: data_source が指定されていればステップを展開
        data_source = sheet_def.get("data_source")
        if data_source:
            steps = _expand_data_source(steps, data_source, input_dir)
        test_cases = [normalize_step(s, j + 1, selector_aliases) for j, s in enumerate(steps)]
        create_sheet(wb, name, test_cases, project_name, is_first=(i == 0))

    if output_path is None:
        output_path = input_path.with_suffix(".xlsx")
    output_path = Path(output_path)
    wb.save(str(output_path))
    print(f"テスト仕様書を生成しました: {output_path}")

    total_steps = sum(len(s.get("steps", [])) for s in sheets)
    print(f"  シート数: {len(sheets)}, 合計ステップ数: {total_steps}")


def main():
    parser = argparse.ArgumentParser(
        description="テスト定義 (.yaml / .txt) → Excel テスト仕様書 生成ツール"
    )
    parser.add_argument("input_file", help="テスト定義ファイル (.yaml / .txt)")
    parser.add_argument("-o", "--output", help="出力先Excelファイルのパス")
    args = parser.parse_args()
    gen_spec(args.input_file, args.output)


if __name__ == "__main__":
    main()
