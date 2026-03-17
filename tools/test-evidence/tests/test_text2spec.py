"""text2spec.py のユニットテスト."""

import sys
from pathlib import Path

import pytest
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from text2spec import parse_text_spec, _parse_selectors_directive


class TestParseSelectorsDirective:
    """@selectors ディレクティブのパースのテスト."""

    def test_single_alias(self):
        result = _parse_selectors_directive("ログインボタン=#login-btn")
        assert result == {"ログインボタン": "#login-btn"}

    def test_multiple_aliases(self):
        result = _parse_selectors_directive("名前=#name, メール=#email")
        assert result == {"名前": "#name", "メール": "#email"}

    def test_spaces_trimmed(self):
        result = _parse_selectors_directive("  名前 = #name ,  メール = #email  ")
        assert result == {"名前": "#name", "メール": "#email"}

    def test_empty_value(self):
        assert _parse_selectors_directive("") == {}

    def test_no_equals(self, capsys):
        result = _parse_selectors_directive("不正な値")
        assert result == {}
        captured = capsys.readouterr()
        assert "警告" in captured.err


class TestParseTextSpec:
    """parse_text_spec のテスト."""

    def test_single_sheet_basic(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text("# ログイン\n「#btn」をクリック\nスクリーンショット\n",
                       encoding="utf-8")
        spec = parse_text_spec(txt)
        assert len(spec["sheets"]) == 1
        assert spec["sheets"][0]["name"] == "ログイン"
        assert len(spec["sheets"][0]["steps"]) == 2
        assert spec["sheets"][0]["steps"][0] == {"item": "「#btn」をクリック"}
        assert spec["sheets"][0]["steps"][1] == {"item": "スクリーンショット"}

    def test_multiple_sheets(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text(
            "# シート1\n「#a」をクリック\n\n# シート2\n「#b」をクリック\n",
            encoding="utf-8")
        spec = parse_text_spec(txt)
        assert len(spec["sheets"]) == 2
        assert spec["sheets"][0]["name"] == "シート1"
        assert spec["sheets"][1]["name"] == "シート2"

    def test_no_heading_default_name(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text("「#btn」をクリック\nスクリーンショット\n",
                       encoding="utf-8")
        spec = parse_text_spec(txt)
        assert spec["sheets"][0]["name"] == "テスト"

    def test_project_directive(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text("@project: テストプロジェクト\n# テスト\n「#x」をクリック\n",
                       encoding="utf-8")
        spec = parse_text_spec(txt)
        assert spec["project"] == "テストプロジェクト"

    def test_base_url_directive(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text("@base_url: http://localhost:5000\n# テスト\n「#x」をクリック\n",
                       encoding="utf-8")
        spec = parse_text_spec(txt)
        assert spec["environment"]["base_url"] == "http://localhost:5000"

    def test_selectors_directive(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text(
            "@selectors: ログインボタン=#login-btn, 検索欄=#search\n"
            "# テスト\nログインボタンをクリック\n",
            encoding="utf-8")
        spec = parse_text_spec(txt)
        assert spec["selectors"] == {"ログインボタン": "#login-btn", "検索欄": "#search"}

    def test_selectors_merge(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text(
            "@selectors: A=#a\n@selectors: B=#b\n# テスト\n「#x」をクリック\n",
            encoding="utf-8")
        spec = parse_text_spec(txt)
        assert spec["selectors"] == {"A": "#a", "B": "#b"}

    def test_data_source_directive(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text(
            "# データテスト\n@data_source: users.csv\n「#name」に「${data:name}」を入力\n",
            encoding="utf-8")
        spec = parse_text_spec(txt)
        assert spec["sheets"][0]["data_source"] == "users.csv"

    def test_comments_ignored(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text(
            "// これはコメント\n# テスト\n// もう一つ\n「#btn」をクリック\n",
            encoding="utf-8")
        spec = parse_text_spec(txt)
        assert len(spec["sheets"][0]["steps"]) == 1

    def test_empty_lines_ignored(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text(
            "\n\n# テスト\n\n「#btn」をクリック\n\n\nスクリーンショット\n\n",
            encoding="utf-8")
        spec = parse_text_spec(txt)
        assert len(spec["sheets"][0]["steps"]) == 2

    def test_empty_file_error(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text("", encoding="utf-8")
        with pytest.raises(SystemExit):
            parse_text_spec(txt)

    def test_comments_only_error(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text("// only comments\n\n// here\n", encoding="utf-8")
        with pytest.raises(SystemExit):
            parse_text_spec(txt)

    def test_utf8_bom(self, tmp_path):
        txt = tmp_path / "test.txt"
        # BOM 付き UTF-8
        txt.write_bytes(b"\xef\xbb\xbf# \xe3\x83\x86\xe3\x82\xb9\xe3\x83\x88\n"
                        b"\xe3\x82\xb9\xe3\x82\xaf\xe3\x83\xaa\xe3\x83\xbc\xe3\x83\xb3"
                        b"\xe3\x82\xb7\xe3\x83\xa7\xe3\x83\x83\xe3\x83\x88\n")
        spec = parse_text_spec(txt)
        assert spec["sheets"][0]["name"] == "テスト"
        assert spec["sheets"][0]["steps"][0]["item"] == "スクリーンショット"

    def test_unknown_directive_warning(self, tmp_path, capsys):
        txt = tmp_path / "test.txt"
        txt.write_text("@unknown: value\n# テスト\n「#x」をクリック\n",
                       encoding="utf-8")
        spec = parse_text_spec(txt)
        captured = capsys.readouterr()
        assert "不明なディレクティブ" in captured.err
        assert len(spec["sheets"]) == 1

    def test_data_source_before_heading(self, tmp_path):
        """@data_source がシート見出しより前にある場合、最初のシートに付与."""
        txt = tmp_path / "test.txt"
        txt.write_text(
            "@data_source: data.csv\n「#x」をクリック\n",
            encoding="utf-8")
        spec = parse_text_spec(txt)
        assert spec["sheets"][0].get("data_source") == "data.csv"

    def test_hash_only_heading(self, tmp_path):
        """'#' のみの見出し行 → デフォルトシート名."""
        txt = tmp_path / "test.txt"
        txt.write_text("#\n「#x」をクリック\n", encoding="utf-8")
        spec = parse_text_spec(txt)
        assert spec["sheets"][0]["name"].startswith("Sheet")


class TestGenSpecTextIntegration:
    """gen_spec.py での .txt 入力統合テスト."""

    def test_txt_to_xlsx(self, tmp_path):
        from gen_spec import gen_spec
        txt = tmp_path / "test.txt"
        txt.write_text(
            "# ログイン\n「http://example.com」に遷移\n「#btn」をクリック\nスクリーンショット\n",
            encoding="utf-8")
        out = tmp_path / "test.xlsx"
        gen_spec(str(txt), str(out))
        assert out.exists()

        from openpyxl import load_workbook
        wb = load_workbook(str(out))
        ws = wb.active
        assert ws.title == "ログイン"
        # navigate
        assert ws["C7"].value == "navigate"
        # click
        assert ws["C8"].value == "click"
        assert ws["D8"].value == "#btn"
        # screenshot
        assert ws["F9"].value == "screenshot"

    def test_txt_with_selectors(self, tmp_path):
        from gen_spec import gen_spec
        txt = tmp_path / "test.txt"
        txt.write_text(
            "@selectors: 送信ボタン=#submit-btn\n"
            "# テスト\n「送信ボタン」をクリック\n",
            encoding="utf-8")
        out = tmp_path / "test.xlsx"
        gen_spec(str(txt), str(out))

        from openpyxl import load_workbook
        wb = load_workbook(str(out))
        ws = wb.active
        assert ws["C7"].value == "click"
        assert ws["D7"].value == "#submit-btn"

    def test_txt_nl_warnings_printed(self, tmp_path, capsys):
        from gen_spec import gen_spec
        txt = tmp_path / "test.txt"
        txt.write_text(
            "# テスト\n「#btnをクリック\n",  # 括弧不整合
            encoding="utf-8")
        out = tmp_path / "test.xlsx"
        gen_spec(str(txt), str(out))
        captured = capsys.readouterr()
        assert "警告" in captured.err
        assert "括弧" in captured.err

    def test_txt_default_output_name(self, tmp_path):
        from gen_spec import gen_spec
        txt = tmp_path / "mytest.txt"
        txt.write_text("# テスト\nスクリーンショット\n", encoding="utf-8")
        gen_spec(str(txt))
        expected_out = tmp_path / "mytest.xlsx"
        assert expected_out.exists()

    def test_txt_multiple_sheets(self, tmp_path):
        from gen_spec import gen_spec
        txt = tmp_path / "test.txt"
        txt.write_text(
            "@project: マルチシートテスト\n"
            "# シート1\n「#a」をクリック\n"
            "# シート2\n「#b」をクリック\n",
            encoding="utf-8")
        out = tmp_path / "test.xlsx"
        gen_spec(str(txt), str(out))

        from openpyxl import load_workbook
        wb = load_workbook(str(out))
        assert len(wb.sheetnames) == 2
        assert wb.sheetnames[0] == "シート1"
        assert wb.sheetnames[1] == "シート2"


class TestSqlExecPrefix:
    """SQL実行/DB実行 プレフィックスのパーステスト."""

    def test_sql_exec_prefix(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text(
            "# DBテスト\nSQL実行: INSERT INTO t VALUES (1, 'x')\n",
            encoding="utf-8")
        spec = parse_text_spec(txt)
        step = spec["sheets"][0]["steps"][0]
        assert step["action"] == "sql_exec"
        assert step["input"] == "INSERT INTO t VALUES (1, 'x')"

    def test_db_exec_prefix(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text(
            "# DBテスト\nDB実行: DELETE FROM t WHERE id = 1\n",
            encoding="utf-8")
        spec = parse_text_spec(txt)
        step = spec["sheets"][0]["steps"][0]
        assert step["action"] == "sql_exec"
        assert step["input"] == "DELETE FROM t WHERE id = 1"

    def test_sql_exec_fullwidth_colon(self, tmp_path):
        txt = tmp_path / "test.txt"
        txt.write_text(
            "# DBテスト\nSQL実行：UPDATE t SET x = 1\n",
            encoding="utf-8")
        spec = parse_text_spec(txt)
        step = spec["sheets"][0]["steps"][0]
        assert step["action"] == "sql_exec"
        assert step["input"] == "UPDATE t SET x = 1"

    def test_sql_exec_item_truncated(self, tmp_path):
        """item フィールドが長いSQLを40文字に切り詰めること."""
        long_sql = "INSERT INTO very_long_table_name (col1, col2, col3) VALUES ('aaa', 'bbb', 'ccc')"
        txt = tmp_path / "test.txt"
        txt.write_text(f"# DBテスト\nSQL実行: {long_sql}\n", encoding="utf-8")
        spec = parse_text_spec(txt)
        step = spec["sheets"][0]["steps"][0]
        assert step["input"] == long_sql
        assert len(step["item"]) <= len("SQL実行: ") + 40

    def test_sql_exec_mixed_with_browser_steps(self, tmp_path):
        """SQL実行とブラウザ操作ステップが混在."""
        txt = tmp_path / "test.txt"
        txt.write_text(
            "# テスト\n"
            "SQL実行: DELETE FROM t\n"
            "「#btn」をクリック\n"
            "DB実行: INSERT INTO t VALUES (1)\n",
            encoding="utf-8")
        spec = parse_text_spec(txt)
        steps = spec["sheets"][0]["steps"]
        assert len(steps) == 3
        assert steps[0]["action"] == "sql_exec"
        assert steps[1] == {"item": "「#btn」をクリック"}
        assert steps[2]["action"] == "sql_exec"
