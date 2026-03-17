"""gen_spec.py のユニットテスト."""

import tempfile
from pathlib import Path

import pytest
import yaml

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from gen_spec import (
    normalize_step, validate_steps, gen_spec, _expand_data_source,
    _interpret_natural_language,
)


class TestNormalizeStep:
    """YAML短縮記法の正規化テスト."""

    def test_basic_action(self):
        step = normalize_step({"action": "click", "selector": "#btn"}, 1)
        assert step["no"] == 1
        assert step["action"] == "click"
        assert step["selector"] == "#btn"
        assert step["item"] == "click: #btn"

    def test_navigate_auto_item(self):
        step = normalize_step({"action": "navigate", "input": "http://example.com"}, 1)
        assert step["action"] == "navigate"
        assert step["input"] == "http://example.com"
        assert step["item"] == "ページを開く"

    def test_include_shorthand(self):
        step = normalize_step({"action": "include", "sheet": "共通ログイン"}, 1)
        assert step["action"] == "include"
        assert step["input"] == "共通ログイン"
        assert "共通ログイン" in step["item"]

    def test_verify_shorthand(self):
        step = normalize_step({
            "verify": "text", "target": "#msg", "expected": "OK"
        }, 1)
        assert step["verify_type"] == "text"
        assert step["verify_target"] == "#msg"
        assert step["expected"] == "OK"

    def test_verify_db_shorthand(self):
        step = normalize_step({
            "verify": "db", "sql": "SELECT COUNT(*) FROM t", "expected": "rowcount:3"
        }, 1)
        assert step["verify_type"] == "db"
        assert step["verify_target"] == "SELECT COUNT(*) FROM t"
        assert step["expected"] == "rowcount:3"

    def test_explicit_item(self):
        step = normalize_step({
            "action": "click", "selector": "#btn", "item": "送信ボタン押下"
        }, 1)
        assert step["item"] == "送信ボタン押下"

    def test_verify_target_fallback(self):
        """verify短縮記法でverify_targetキーを使った場合のフォールバック."""
        step = normalize_step({
            "verify": "text", "verify_target": "#welcome-msg", "expected": "OK"
        }, 1)
        assert step["verify_target"] == "#welcome-msg"

    def test_longform_verify(self):
        """verify短縮記法を使わず、verify_type/verify_target で指定."""
        step = normalize_step({
            "verify_type": "visible", "verify_target": "#dashboard"
        }, 1)
        assert step["verify_type"] == "visible"
        assert step["verify_target"] == "#dashboard"


class TestValidateSteps:
    """テストステップのバリデーションテスト."""

    def test_valid_steps(self):
        steps = [
            {"action": "navigate", "input": "http://example.com"},
            {"action": "click", "selector": "#btn"},
            {"verify": "text", "target": "#msg"},
        ]
        errors = validate_steps("Sheet1", steps)
        assert errors == []

    def test_invalid_action(self):
        errors = validate_steps("Sheet1", [{"action": "invalid_action"}])
        assert len(errors) == 1
        assert "無効な操作種別" in errors[0]

    def test_invalid_verify_type(self):
        errors = validate_steps("Sheet1", [{"verify": "invalid_verify"}])
        assert len(errors) == 1
        assert "無効な検証種別" in errors[0]

    def test_click_without_selector(self):
        errors = validate_steps("Sheet1", [{"action": "click"}])
        assert len(errors) == 1
        assert "セレクタが未指定" in errors[0]

    def test_navigate_without_input(self):
        errors = validate_steps("Sheet1", [{"action": "navigate"}])
        assert len(errors) == 1
        assert "URLが未指定" in errors[0]

    def test_include_without_sheet(self):
        errors = validate_steps("Sheet1", [{"action": "include"}])
        assert len(errors) == 1
        assert "シート名が未指定" in errors[0]

    def test_include_with_sheet_is_valid(self):
        errors = validate_steps("Sheet1", [
            {"action": "include", "sheet": "共通手順"}
        ])
        assert errors == []


class TestGenSpec:
    """gen_spec関数のE2Eテスト."""

    def test_generates_xlsx(self, tmp_path):
        yaml_file = tmp_path / "test.yaml"
        yaml_file.write_text(yaml.dump({
            "project": "テスト",
            "sheets": [{
                "name": "Sheet1",
                "steps": [
                    {"action": "navigate", "input": "http://example.com"},
                    {"action": "click", "selector": "#btn", "item": "ボタン押下"},
                ],
            }],
        }, allow_unicode=True), encoding="utf-8")

        output_file = tmp_path / "test.xlsx"
        gen_spec(str(yaml_file), str(output_file))
        assert output_file.exists()
        assert output_file.stat().st_size > 0

    def test_validation_error_exits(self, tmp_path):
        yaml_file = tmp_path / "bad.yaml"
        yaml_file.write_text(yaml.dump({
            "sheets": [{
                "name": "Sheet1",
                "steps": [{"action": "invalid_action"}],
            }],
        }, allow_unicode=True), encoding="utf-8")

        with pytest.raises(SystemExit):
            gen_spec(str(yaml_file))

    def test_missing_file_exits(self):
        with pytest.raises(SystemExit):
            gen_spec("/nonexistent/file.yaml")

    def test_data_driven_csv(self, tmp_path):
        """data_source でCSVを指定するとステップが展開される."""
        csv_file = tmp_path / "users.csv"
        csv_file.write_text("name,email\n田中,tanaka@test.com\n鈴木,suzuki@test.com\n",
                            encoding="utf-8")
        yaml_file = tmp_path / "test.yaml"
        yaml_file.write_text(yaml.dump({
            "project": "DDT",
            "sheets": [{
                "name": "ユーザー登録",
                "data_source": "users.csv",
                "steps": [
                    {"action": "input", "selector": "#name", "input": "${data:name}",
                     "item": "名前入力"},
                    {"action": "input", "selector": "#email", "input": "${data:email}",
                     "item": "メール入力"},
                ],
            }],
        }, allow_unicode=True), encoding="utf-8")

        output_file = tmp_path / "test.xlsx"
        gen_spec(str(yaml_file), str(output_file))
        assert output_file.exists()

        from openpyxl import load_workbook
        wb = load_workbook(str(output_file))
        ws = wb.active
        # 2名 × 2ステップ = 4行のデータ（7行目から開始）
        assert ws["A7"].value == 1
        assert ws["A10"].value == 4
        # ${data:name} が展開されている
        assert "田中" in (ws["E7"].value or "")
        assert "鈴木" in (ws["E9"].value or "")

    def test_data_driven_json(self, tmp_path):
        """data_source でJSONを指定するとステップが展開される."""
        import json
        json_file = tmp_path / "items.json"
        json_file.write_text(json.dumps([
            {"product": "商品A", "price": "1000"},
            {"product": "商品B", "price": "2000"},
            {"product": "商品C", "price": "3000"},
        ], ensure_ascii=False), encoding="utf-8")
        yaml_file = tmp_path / "test.yaml"
        yaml_file.write_text(yaml.dump({
            "project": "DDT",
            "sheets": [{
                "name": "商品登録",
                "data_source": "items.json",
                "steps": [
                    {"action": "input", "selector": "#product", "input": "${data:product}",
                     "item": "商品名入力"},
                ],
            }],
        }, allow_unicode=True), encoding="utf-8")

        output_file = tmp_path / "test.xlsx"
        gen_spec(str(yaml_file), str(output_file))

        from openpyxl import load_workbook
        wb = load_workbook(str(output_file))
        ws = wb.active
        # 3件 × 1ステップ = 3行
        assert ws["A9"].value == 3
        assert "商品C" in (ws["E9"].value or "")


class TestNormalizeStepExtended:
    """新アクションの正規化テスト."""

    def test_capture(self):
        step = normalize_step({"action": "capture", "selector": "#id", "input": "order_id"}, 1)
        assert step["action"] == "capture"
        assert step["input"] == "order_id"

    def test_new_tab(self):
        step = normalize_step({"action": "new_tab", "selector": "#link", "input": "report"}, 1)
        assert step["action"] == "new_tab"

    def test_switch_tab(self):
        step = normalize_step({"action": "switch_tab", "input": "main"}, 1)
        assert step["action"] == "switch_tab"

    def test_download(self):
        step = normalize_step({"action": "download", "selector": "#export"}, 1)
        assert step["action"] == "download"

    def test_if_ok(self):
        step = normalize_step({"action": "if_ok", "input": "click", "selector": "#next"}, 1)
        assert step["action"] == "if_ok"
        assert step["input"] == "click"

    def test_if_ng(self):
        step = normalize_step({"action": "if_ng", "input": "click", "selector": "#retry"}, 1)
        assert step["action"] == "if_ng"

    def test_retry_in_note(self):
        step = normalize_step({"action": "click", "selector": "#btn", "retry": 3}, 1)
        assert "retry:3" in step["note"]

    def test_retry_and_wait_in_note(self):
        step = normalize_step({"action": "click", "selector": "#btn",
                               "retry": 2, "retry_wait": 500}, 1)
        assert "retry:2" in step["note"]
        assert "retry_wait:500" in step["note"]

    def test_verify_download(self):
        step = normalize_step({"verify": "download", "target": "export.csv",
                               "expected": "contains:田中"}, 1)
        assert step["verify_type"] == "download"


class TestValidateStepsExtended:
    """新アクションのバリデーションテスト."""

    def test_capture_valid(self):
        errors = validate_steps("S", [{"action": "capture", "selector": "#id", "input": "var"}])
        assert errors == []

    def test_capture_no_selector(self):
        errors = validate_steps("S", [{"action": "capture", "input": "var"}])
        assert any("セレクタが未指定" in e for e in errors)

    def test_capture_no_input(self):
        errors = validate_steps("S", [{"action": "capture", "selector": "#id"}])
        assert any("変数名" in e for e in errors)

    def test_switch_tab_no_input(self):
        errors = validate_steps("S", [{"action": "switch_tab"}])
        assert any("タブ名" in e for e in errors)

    def test_download_no_selector(self):
        errors = validate_steps("S", [{"action": "download"}])
        assert any("セレクタが未指定" in e for e in errors)

    def test_new_actions_valid(self):
        steps = [
            {"action": "capture", "selector": "#x", "input": "v"},
            {"action": "new_tab", "selector": "#link", "input": "tab1"},
            {"action": "switch_tab", "input": "tab1"},
            {"action": "close_tab", "input": "tab1"},
            {"action": "download", "selector": "#dl"},
            {"action": "if_ok", "input": "click", "selector": "#btn"},
            {"action": "if_ng", "input": "click", "selector": "#retry"},
        ]
        errors = validate_steps("S", steps)
        assert errors == []


class TestExpandDataSource:
    """data_source 展開テスト."""

    def test_csv_expansion(self, tmp_path):
        csv_file = tmp_path / "data.csv"
        csv_file.write_text("name,email\nAlice,a@b.com\nBob,b@c.com\n",
                            encoding="utf-8")
        steps = [
            {"action": "input", "selector": "#name", "input": "${data:name}", "item": "名前"},
            {"action": "input", "selector": "#email", "input": "${data:email}", "item": "メール"},
        ]
        expanded = _expand_data_source(steps, "data.csv", tmp_path)
        assert len(expanded) == 4  # 2行 × 2ステップ
        assert expanded[0]["input"] == "Alice"
        assert expanded[1]["input"] == "a@b.com"
        assert expanded[2]["input"] == "Bob"
        assert expanded[3]["input"] == "b@c.com"
        # テスト項目名にrow番号が付与
        assert "[row1:" in expanded[0]["item"]
        assert "[row2:" in expanded[2]["item"]

    def test_json_expansion(self, tmp_path):
        import json as json_mod
        json_file = tmp_path / "data.json"
        json_file.write_text(json_mod.dumps([
            {"x": "1"}, {"x": "2"}, {"x": "3"}
        ]), encoding="utf-8")
        steps = [{"action": "input", "selector": "#x", "input": "${data:x}", "item": "入力"}]
        expanded = _expand_data_source(steps, "data.json", tmp_path)
        assert len(expanded) == 3
        assert expanded[0]["input"] == "1"
        assert expanded[2]["input"] == "3"

    def test_missing_file_exits(self, tmp_path):
        with pytest.raises(SystemExit):
            _expand_data_source([], "missing.csv", tmp_path)

    def test_unresolved_var_unchanged(self, tmp_path):
        csv_file = tmp_path / "data.csv"
        csv_file.write_text("a\n1\n", encoding="utf-8")
        steps = [{"action": "input", "selector": "#x", "input": "${data:missing}", "item": "test"}]
        expanded = _expand_data_source(steps, "data.csv", tmp_path)
        assert expanded[0]["input"] == "${data:missing}"


class TestNaturalLanguage:
    """自然言語テスト記述の解釈テスト."""

    # --- 操作パターン ---

    def test_click_selector(self):
        result = _interpret_natural_language({"item": "「#btn」をクリック"})
        assert result["action"] == "click"
        assert result["selector"] == "#btn"

    def test_click_button_name(self):
        result = _interpret_natural_language({"item": "ログインボタンをクリック"})
        assert result["action"] == "click"
        assert result["selector"] == "ログイン"

    def test_click_link(self):
        result = _interpret_natural_language({"item": "詳細リンクをクリック"})
        assert result["action"] == "click"
        assert result["selector"] == "詳細"

    def test_input_pattern(self):
        result = _interpret_natural_language({"item": "「#name」に「田中」を入力"})
        assert result["action"] == "input"
        assert result["selector"] == "#name"
        assert result["input"] == "田中"

    def test_input_field_name(self):
        result = _interpret_natural_language({"item": "ユーザー名欄に「admin」を入力"})
        assert result["action"] == "input"
        assert result["selector"] == "ユーザー名"
        assert result["input"] == "admin"

    def test_select_pattern(self):
        result = _interpret_natural_language({"item": "「#priority」で「高」を選択"})
        assert result["action"] == "select"
        assert result["selector"] == "#priority"
        assert result["input"] == "高"

    def test_navigate_pattern(self):
        result = _interpret_natural_language({"item": "${base_url}/login を開く"})
        assert result["action"] == "navigate"
        assert result["input"] == "${base_url}/login "

    def test_navigate_transition(self):
        result = _interpret_natural_language({"item": "「http://example.com」に遷移"})
        assert result["action"] == "navigate"
        assert result["input"] == "http://example.com"

    def test_wait_seconds(self):
        result = _interpret_natural_language({"item": "3秒待機"})
        assert result["action"] == "wait"
        assert result["input"] == "3000"

    def test_wait_for_element(self):
        result = _interpret_natural_language({"item": "「#dashboard」が表示されるまで待機"})
        assert result["action"] == "wait_for"
        assert result["selector"] == "#dashboard"

    def test_hover_pattern(self):
        result = _interpret_natural_language({"item": "「.menu-item」にマウスオーバー"})
        assert result["action"] == "hover"
        assert result["selector"] == ".menu-item"

    def test_keyboard_enter(self):
        result = _interpret_natural_language({"item": "Enterキーを押す"})
        assert result["action"] == "keyboard"
        assert result["input"] == "Enter"

    def test_capture_pattern(self):
        result = _interpret_natural_language({"item": "「#order-id」の値を保存"})
        assert result["action"] == "capture"
        assert result["selector"] == "#order-id"

    def test_upload_pattern(self):
        result = _interpret_natural_language({"item": "「report.pdf」をアップロード"})
        assert result["action"] == "upload"
        assert result["input"] == "report.pdf"

    def test_scroll_pattern(self):
        result = _interpret_natural_language({"item": "「#footer」までスクロール"})
        assert result["action"] == "scroll"
        assert result["selector"] == "#footer"

    # --- 検証パターン ---

    def test_verify_text(self):
        result = _interpret_natural_language({"item": "「#msg」に「登録完了」と表示"})
        assert result["verify"] == "text"
        assert result["target"] == "#msg"
        assert result["expected"] == "登録完了"

    def test_verify_contains(self):
        result = _interpret_natural_language({"item": "「#msg」に「成功」を含む"})
        assert result["verify"] == "text"
        assert result["target"] == "#msg"
        assert result["expected"] == "contains:成功"

    def test_verify_visible(self):
        result = _interpret_natural_language({"item": "「#dashboard」が表示されること"})
        assert result["verify"] == "visible"
        assert result["target"] == "#dashboard"

    def test_verify_visible_short(self):
        result = _interpret_natural_language({"item": "「#dashboard」が表示される"})
        assert result["verify"] == "visible"

    def test_verify_hidden(self):
        result = _interpret_natural_language({"item": "「#error-msg」が非表示"})
        assert result["verify"] == "hidden"
        assert result["target"] == "#error-msg"

    def test_verify_url(self):
        result = _interpret_natural_language({"item": "URLが「/dashboard」であること"})
        assert result["verify"] == "url"
        assert result["expected"] == "/dashboard"

    def test_verify_screenshot(self):
        result = _interpret_natural_language({"item": "スクリーンショットを取得"})
        assert result["verify"] == "screenshot"

    def test_verify_screenshot_short(self):
        result = _interpret_natural_language({"item": "スクリーンショット"})
        assert result["verify"] == "screenshot"

    # --- スキップ/優先ルール ---

    def test_explicit_action_skips_nl(self):
        """action が既にあれば NL 解釈しない."""
        result = _interpret_natural_language({"action": "click", "item": "何かの説明"})
        assert result["action"] == "click"
        assert "selector" not in result

    def test_explicit_verify_skips_nl(self):
        """verify が既にあれば NL 解釈しない."""
        result = _interpret_natural_language({"verify": "visible", "item": "何かの説明"})
        assert result["verify"] == "visible"

    def test_no_match_unchanged(self):
        """パターン不一致の場合は変更なし."""
        result = _interpret_natural_language({"item": "特に何もしない備考"})
        assert "action" not in result
        assert "verify" not in result

    def test_empty_item(self):
        result = _interpret_natural_language({"item": ""})
        assert "action" not in result

    def test_no_item(self):
        result = _interpret_natural_language({})
        assert "action" not in result

    # --- セレクタ辞書 ---

    def test_selector_alias_click(self):
        aliases = {"ログインボタン": "#login-btn"}
        result = _interpret_natural_language(
            {"item": "「ログインボタン」をクリック"}, aliases)
        assert result["action"] == "click"
        assert result["selector"] == "#login-btn"

    def test_selector_alias_verify(self):
        aliases = {"成功メッセージ": "#msg.success"}
        result = _interpret_natural_language(
            {"item": "「成功メッセージ」が表示されること"}, aliases)
        assert result["verify"] == "visible"
        assert result["target"] == "#msg.success"

    def test_selector_alias_button_suffix(self):
        """〇〇ボタンをクリック でセレクタ辞書を参照."""
        aliases = {"ログイン": "#login-btn"}
        result = _interpret_natural_language(
            {"item": "ログインボタンをクリック"}, aliases)
        assert result["selector"] == "#login-btn"

    def test_selector_alias_not_found(self):
        """辞書にない名前はそのまま保持."""
        result = _interpret_natural_language(
            {"item": "「未定義要素」をクリック"}, {"別の要素": "#x"})
        assert result["selector"] == "未定義要素"

    # --- normalize_step との統合 ---

    def test_normalize_step_with_nl(self):
        """normalize_step が NL 解釈を適用する."""
        step = normalize_step({"item": "「#btn」をクリック"}, 1)
        assert step["action"] == "click"
        assert step["selector"] == "#btn"
        assert step["no"] == 1

    def test_normalize_step_nl_with_aliases(self):
        aliases = {"送信ボタン": "#submit-btn"}
        step = normalize_step({"item": "「送信ボタン」をクリック"}, 1, aliases)
        assert step["selector"] == "#submit-btn"

    def test_gen_spec_nl_yaml(self, tmp_path):
        """gen_spec が NL記法の YAML を正しく処理する."""
        yaml_file = tmp_path / "nl_test.yaml"
        yaml_file.write_text(yaml.dump({
            "project": "NLテスト",
            "selectors": {"ログインボタン": "#login-btn"},
            "sheets": [{
                "name": "NLシート",
                "steps": [
                    {"item": "「http://example.com」に遷移"},
                    {"item": "「#username」に「admin」を入力"},
                    {"item": "「ログインボタン」をクリック"},
                    {"item": "「#dashboard」が表示されること"},
                    {"item": "スクリーンショットを取得"},
                ],
            }],
        }, allow_unicode=True), encoding="utf-8")

        output_file = tmp_path / "nl_test.xlsx"
        gen_spec(str(yaml_file), str(output_file))
        assert output_file.exists()

        from openpyxl import load_workbook
        wb = load_workbook(str(output_file))
        ws = wb.active
        # navigate が生成されている
        assert ws["C7"].value == "navigate"
        # input が生成されている
        assert ws["C8"].value == "input"
        assert ws["D8"].value == "#username"
        assert ws["E8"].value == "admin"
        # click with alias
        assert ws["C9"].value == "click"
        assert ws["D9"].value == "#login-btn"
        # verify visible
        assert ws["F10"].value == "visible"
        assert ws["G10"].value == "#dashboard"
        # screenshot
        assert ws["F11"].value == "screenshot"
