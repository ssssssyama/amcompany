"""デモ用テスト仕様書を生成するスクリプト."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(name="Yu Gothic", size=10, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Yu Gothic", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def create_sheet(wb, sheet_name, test_cases, is_first=False):
    if is_first:
        ws = wb.active
        ws.title = sheet_name
    else:
        ws = wb.create_sheet(sheet_name)

    # ヘッダー情報
    ws["A1"] = "テスト仕様書"
    ws["A1"].font = Font(name="Yu Gothic", size=14, bold=True)
    ws["A2"] = "プロジェクト名:"
    ws["B2"] = "デモアプリ - タスク管理"
    ws["A3"] = "テスト日:"
    ws["A4"] = "テスト担当者:"
    ws["B4"] = "自動テスト"

    headers = [
        ("A", "No.", 6), ("B", "テスト項目", 30), ("C", "操作種別", 12),
        ("D", "対象セレクタ", 30), ("E", "入力値", 25),
        ("F", "検証種別", 12), ("G", "検証対象", 35),
        ("H", "期待値", 30), ("I", "結果", 8),
        ("J", "エビデンス", 30), ("K", "備考", 25),
    ]
    for col, title, width in headers:
        cell = ws[f"{col}6"]
        cell.value = title
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        ws.column_dimensions[col].width = width

    for i, tc in enumerate(test_cases):
        row = 7 + i
        values = [
            tc.get("no", i + 1), tc.get("item", ""),
            tc.get("action", ""), tc.get("selector", ""),
            tc.get("input", ""), tc.get("verify_type", ""),
            tc.get("verify_target", ""), tc.get("expected", ""),
            "", "", tc.get("note", ""),
        ]
        for j, val in enumerate(values):
            col = chr(65 + j)
            cell = ws[f"{col}{row}"]
            cell.value = val
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            cell.alignment = LEFT if j > 0 else CENTER
        ws.row_dimensions[row].height = 80

    return ws


def main():
    wb = Workbook()

    # --- シート1: ログイン機能 ---
    login_cases = [
        {
            "no": 1, "item": "ログインページを開く",
            "action": "navigate", "input": "${base_url}",
            "verify_type": "screenshot",
            "note": "ログイン画面が表示されること",
        },
        {
            "no": 2, "item": "ユーザー名を入力",
            "action": "input", "selector": "#username",
            "input": "${test_user}",
        },
        {
            "no": 3, "item": "パスワードを入力",
            "action": "input", "selector": "#password",
            "input": "${test_pass}",
        },
        {
            "no": 4, "item": "入力状態のスクリーンショット",
            "action": "wait", "input": "500",
            "verify_type": "screenshot",
            "note": "ユーザー名・パスワードが入力されていること",
        },
        {
            "no": 5, "item": "ログインボタンを押下",
            "action": "click", "selector": "#login-button",
        },
        {
            "no": 6, "item": "ダッシュボード表示を待機",
            "action": "wait_for", "selector": "#welcome-msg",
            "input": "10000",
            "verify_type": "text",
            "verify_target": "#welcome-msg",
            "expected": "contains:テストユーザー",
            "note": "wait_for: 非同期レンダリング後にウェルカムメッセージが表示されること",
        },
        {
            "no": 7, "item": "ダッシュボード全体のスクリーンショット",
            "action": "wait", "input": "500",
            "verify_type": "screenshot",
            "note": "ダッシュボードが正しく表示されること",
        },
        {
            "no": 8, "item": "タスク一覧が表示されていること",
            "action": "",
            "verify_type": "visible",
            "verify_target": "#task-tbody",
            "note": "初期タスクが表示されていること",
        },
        {
            "no": 9, "item": "タスク件数の確認",
            "action": "",
            "verify_type": "text",
            "verify_target": "#task-count",
            "expected": "contains:3件",
            "note": "初期タスクが3件であること",
        },
        {
            "no": 10, "item": "DB: ログイン履歴が記録されていること",
            "action": "",
            "verify_type": "db",
            "verify_target": "SELECT COUNT(*) FROM login_history WHERE user_id = 2 AND success = 1",
            "expected": "contains:1",
            "note": "testuser(id=2)のログイン成功が1件記録されていること",
        },
        {
            "no": 11, "item": "DB: ユーザー情報の確認",
            "action": "",
            "verify_type": "db",
            "verify_target": "SELECT display_name FROM users WHERE username = 'testuser'",
            "expected": "テストユーザー",
            "note": "testuserの表示名がDBと一致すること",
        },
    ]
    create_sheet(wb, "ログイン機能", login_cases, is_first=True)

    # --- シート2: タスク管理機能 ---
    task_cases = [
        {
            "no": 1, "item": "ログインページを開く",
            "action": "navigate", "input": "${base_url}",
            "verify_type": "screenshot",
            "note": "前提: ログインページを開く",
        },
        {
            "no": 2, "item": "ユーザー名入力",
            "action": "input", "selector": "#username",
            "input": "${test_user}",
        },
        {
            "no": 3, "item": "パスワード入力",
            "action": "input", "selector": "#password",
            "input": "${test_pass}",
        },
        {
            "no": 4, "item": "ログイン実行",
            "action": "click", "selector": "#login-button",
        },
        {
            "no": 5, "item": "ダッシュボード表示を待機",
            "action": "wait_for", "selector": "#dashboard",
            "input": "10000",
            "verify_type": "visible",
            "verify_target": "#dashboard",
            "note": "wait_for: ダッシュボードが表示されるまで待機",
        },
        {
            "no": 6, "item": "DB: 初期タスク件数（rowcount検証）",
            "action": "",
            "verify_type": "db",
            "verify_target": "SELECT * FROM tasks",
            "expected": "rowcount:3",
            "note": "rowcount: 初期タスクが3行あること",
        },
        {
            "no": 7, "item": "DB: 完了タスクが存在する（not_empty検証）",
            "action": "",
            "verify_type": "db",
            "verify_target": "SELECT * FROM tasks WHERE status = '完了'",
            "expected": "not_empty",
            "note": "not_empty: 完了タスクが1件以上あること",
        },
        {
            "no": 8, "item": "DB: 削除済みタスクは無い（empty検証）",
            "action": "",
            "verify_type": "db",
            "verify_target": "SELECT * FROM tasks WHERE title = '存在しないタスク'",
            "expected": "empty",
            "note": "empty: 存在しない条件の結果が0行であること",
        },
        {
            "no": 9, "item": "タスク名を入力",
            "action": "input", "selector": "#task-name",
            "input": "テスト自動化の検討",
        },
        {
            "no": 10, "item": "優先度を「高」に変更",
            "action": "select", "selector": "#task-priority",
            "input": "高",
        },
        {
            "no": 11, "item": "タスク追加ボタンを押下",
            "action": "click", "selector": "#add-task-button",
            "verify_type": "text",
            "verify_target": "#task-count",
            "expected": "contains:4件",
            "note": "タスクが4件に増えること",
        },
        {
            "no": 12, "item": "DB: タスク追加後の行数（rowcount検証）",
            "action": "",
            "verify_type": "db",
            "verify_target": "SELECT * FROM tasks",
            "expected": "rowcount:4",
            "note": "rowcount: 追加後にタスクが4行になっていること",
        },
        {
            "no": 13, "item": "DB: タスクが1件以上ある（rowcount>=検証）",
            "action": "",
            "verify_type": "db",
            "verify_target": "SELECT * FROM tasks WHERE priority = '高'",
            "expected": "rowcount:>=1",
            "note": "rowcount:>=N: 優先度「高」のタスクが1件以上あること",
        },
        {
            "no": 14, "item": "DB: 追加タスクの複数カラム検証（rows検証）",
            "action": "",
            "verify_type": "db",
            "verify_target": "SELECT title, priority, status FROM tasks WHERE id = 4",
            "expected": "rows:title=テスト自動化の検討,priority=高,status=未着手",
            "note": "rows: 1行の複数カラムを同時検証",
        },
        {
            "no": 15, "item": "追加後のスクリーンショット",
            "action": "wait", "input": "500",
            "verify_type": "screenshot",
            "note": "新しいタスクが一覧に表示されていること",
        },
        {
            "no": 16, "item": "ログアウトボタンを押下",
            "action": "click", "selector": "#logout-button",
            "verify_type": "visible",
            "verify_target": "#login-page",
            "note": "ログインページに戻ること",
        },
        {
            "no": 17, "item": "ログアウト後のスクリーンショット",
            "action": "wait", "input": "500",
            "verify_type": "screenshot",
            "note": "ログインページが表示されていること",
        },
    ]
    create_sheet(wb, "タスク管理機能", task_cases)

    wb.save("demo_test_spec.xlsx")
    print("デモ用テスト仕様書を作成しました: demo_test_spec.xlsx")


if __name__ == "__main__":
    main()
