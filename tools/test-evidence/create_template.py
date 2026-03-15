"""テスト仕様書のExcelテンプレートを生成するスクリプト."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(name="Yu Gothic", size=10, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Yu Gothic", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def create_template(output_path: str = "examples/test_spec_template.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "テスト仕様書"

    # ヘッダー情報
    ws["A1"] = "テスト仕様書"
    ws["A1"].font = Font(name="Yu Gothic", size=14, bold=True)
    ws["A2"] = "プロジェクト名:"
    ws["B2"] = "サンプルプロジェクト"
    ws["A3"] = "テスト日:"
    ws["A4"] = "テスト担当者:"

    # テストケーステーブルのヘッダー（6行目）
    headers = [
        ("A", "No.", 6),
        ("B", "テスト項目", 30),
        ("C", "操作種別", 12),
        ("D", "対象セレクタ", 25),
        ("E", "入力値/期待値", 20),
        ("F", "検証種別", 12),
        ("G", "検証対象", 25),
        ("H", "期待値", 20),
        ("I", "結果", 8),
        ("J", "エビデンス", 30),
        ("K", "備考", 20),
    ]

    for col, title, width in headers:
        cell = ws[f"{col}6"]
        cell.value = title
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        ws.column_dimensions[col].width = width

    # サンプルテストケース
    test_cases = [
        {
            "no": 1,
            "item": "ログインページ表示",
            "action": "navigate",
            "selector": "",
            "input": "https://example.com/login",
            "verify_type": "screenshot",
            "verify_target": "",
            "expected": "",
            "note": "ログインページが表示されること",
        },
        {
            "no": 2,
            "item": "ユーザー名入力",
            "action": "input",
            "selector": "#username",
            "input": "testuser",
            "verify_type": "",
            "verify_target": "",
            "expected": "",
            "note": "",
        },
        {
            "no": 3,
            "item": "パスワード入力",
            "action": "input",
            "selector": "#password",
            "input": "password123",
            "verify_type": "",
            "verify_target": "",
            "expected": "",
            "note": "",
        },
        {
            "no": 4,
            "item": "ログインボタン押下",
            "action": "click",
            "selector": "#login-button",
            "input": "",
            "verify_type": "text",
            "verify_target": ".welcome-message",
            "expected": "ようこそ、testuser さん",
            "note": "ダッシュボードに遷移すること",
        },
        {
            "no": 5,
            "item": "ダッシュボード表示確認",
            "action": "wait",
            "selector": "",
            "input": "1000",
            "verify_type": "screenshot",
            "verify_target": "",
            "expected": "",
            "note": "ダッシュボードが正しく表示されること",
        },
        {
            "no": 6,
            "item": "DB: ログイン履歴確認",
            "action": "",
            "selector": "",
            "input": "",
            "verify_type": "db",
            "verify_target": "SELECT COUNT(*) FROM login_history WHERE username='testuser'",
            "expected": "1",
            "note": "ログイン履歴がDBに記録されること",
        },
    ]

    for i, tc in enumerate(test_cases):
        row = 7 + i
        values = [
            tc["no"], tc["item"], tc["action"], tc["selector"],
            tc["input"], tc["verify_type"], tc["verify_target"],
            tc["expected"], "", "", tc["note"],
        ]
        for j, val in enumerate(values):
            col = chr(65 + j)  # A, B, C, ...
            cell = ws[f"{col}{row}"]
            cell.value = val
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            cell.alignment = LEFT if j > 0 else CENTER

    # エビデンス列の行の高さを確保
    for row in range(7, 7 + len(test_cases)):
        ws.row_dimensions[row].height = 80

    # 操作種別の凡例シート
    legend = wb.create_sheet("凡例")
    legend["A1"] = "操作種別"
    legend["A1"].font = Font(name="Yu Gothic", size=12, bold=True)

    actions = [
        ("navigate", "指定URLに遷移する"),
        ("click", "指定セレクタの要素をクリック"),
        ("input", "指定セレクタにテキストを入力"),
        ("select", "指定セレクタのドロップダウンで値を選択"),
        ("wait", "指定ミリ秒待機する"),
    ]
    for i, (action, desc) in enumerate(actions):
        legend[f"A{i + 3}"] = action
        legend[f"B{i + 3}"] = desc

    legend["A9"] = "検証種別"
    legend["A9"].font = Font(name="Yu Gothic", size=12, bold=True)

    verifications = [
        ("screenshot", "スクリーンショットを撮影してエビデンスに貼付"),
        ("text", "指定セレクタのテキストが期待値と一致するか検証"),
        ("value", "指定セレクタのvalue属性が期待値と一致するか検証"),
        ("visible", "指定セレクタの要素が表示されているか検証"),
        ("url", "現在のURLが期待値と一致するか検証"),
        ("db", "SQLクエリの結果が期待値と一致するか検証"),
    ]
    for i, (vtype, desc) in enumerate(verifications):
        legend[f"A{i + 11}"] = vtype
        legend[f"B{i + 11}"] = desc

    wb.save(output_path)
    print(f"テンプレートを作成しました: {output_path}")


if __name__ == "__main__":
    create_template()
